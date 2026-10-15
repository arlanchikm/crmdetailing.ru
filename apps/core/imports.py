"""Owner-reviewed CSV/Excel imports with a reversible, conflict-checked journal."""
import csv
import hashlib
import io
import json
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django import forms
from django.core.exceptions import ValidationError
from django.core.serializers.json import DjangoJSONEncoder
from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from apps.customers.models import Client, Car, VehicleClass, normalize_phone, normalize_plate
from apps.inventory.models import Consumable, InventoryCategory, InventoryMovement
from apps.staff.models import Employee, CompanyRole
from apps.tenants.models import Tenant, Subscription
from apps.tenants.plans import counted_employees, active_employee_limit
from .models import ImportBatch, log_event
from .phones import format_phone

MAX_ROWS=5000
KINDS={'clients':'Клиенты','cars':'Автомобили','employees':'Сотрудники','consumables':'Расходники и остатки'}
FIELDS={
 'clients':[('name','Имя',True),('last_name','Фамилия',False),('phone','Телефон',True),('email','Email',False),('notes','Заметки',False)],
 'cars':[('client_phone','Телефон клиента',True),('plate','Госномер',False),('vin','VIN',False),('vehicle_class','Класс автомобиля',False),('make','Марка',False),('model','Модель',False),('color','Цвет',False),('notes','Заметки',False)],
 'employees':[('name','Имя',True),('email','Email пользователя CRM',False),('company_role','Роль',False),('active','Активен',False),('salary_mode','Схема зарплаты',False),('salary_percent','Процент',False),('fixed_monthly_salary','Ставка за месяц',False)],
 'consumables':[('name','Название',True),('category','Категория',True),('unit','Единица',False),('cost_per_unit','Стоимость за единицу',True),('planned_stock','Плановый остаток',False),('actual_stock','Фактический остаток',False),('minimum_stock','Минимальный остаток',False),('active','Активен',False)]}
MODELS={'clients':Client,'cars':Car,'employees':Employee,'consumables':Consumable,'category':InventoryCategory,'movement':InventoryMovement}
ALIASES={'phone':['номер телефона','телефон клиента','phone number','mobile','телефон *'],
 'client_phone':['телефон','phone','client phone'], 'company_role':['должность','role','position'],
 'cost_per_unit':['цена','стоимость','price','cost'], 'planned_stock':['план','остаток','stock'],
 'actual_stock':['факт','фактический остаток'], 'fixed_monthly_salary':['ставка','оклад'], 'unit':['единица измерения','ед изм'], 'plate':['гос номер','номер авто','license plate'], 'name':['клиент','сотрудник','first name','название']}


def norm(value): return re.sub(r'[^\w]', '', str(value).casefold())


def read_table(upload, sheet=''):
    if upload.size>10*1024*1024: raise ValueError('Максимальный размер файла — 10 МБ.')
    raw=upload.read();extension=Path(upload.name).suffix.lower()
    try:
        if extension=='.csv':
            if raw.startswith((b'\xff\xfe',b'\xfe\xff')): text=raw.decode('utf-16')
            else:
                try: text=raw.decode('utf-8-sig')
                except UnicodeDecodeError: text=raw.decode('cp1251')
            try: dialect=csv.Sniffer().sniff(text[:10000],delimiters=',;\t')
            except csv.Error: dialect=csv.excel
            iterator=csv.reader(io.StringIO(text),dialect)
        elif extension=='.xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist())>50*1024*1024:
                    raise ValueError('Распакованный Excel превышает 50 МБ. Разделите файл.')
            workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
            if sheet and sheet not in workbook.sheetnames: raise ValueError('Лист с таким именем не найден.')
            worksheet=workbook[sheet] if sheet else workbook.worksheets[0]
            iterator=worksheet.iter_rows(values_only=True)
        elif extension=='.xls':
            import xlrd
            workbook=xlrd.open_workbook(file_contents=raw,on_demand=True)
            worksheet=workbook.sheet_by_name(sheet) if sheet else workbook.sheet_by_index(0)
            iterator=(worksheet.row_values(i) for i in range(worksheet.nrows))
        else: raise ValueError('Загрузите CSV, XLSX или XLS.')
        header=next(iterator,None)
        if not header or len(header)>100: raise ValueError('Первая строка должна содержать до 100 названий колонок.')
        columns=[str(v or '').strip() or f'Колонка {i+1}' for i,v in enumerate(header)]
        rows=[]
        for number,row in enumerate(iterator,2):
            if number > 20000: raise ValueError('Слишком много пустых строк. Удалите лишние строки в конце таблицы.')
            values=[str(int(v)) if isinstance(v,float) and v.is_integer() else str(v) if v is not None else '' for v in row[:len(columns)]]
            values += ['']*(len(columns)-len(values))
            if not any(v.strip() for v in values): continue
            if len(rows)>=MAX_ROWS: raise ValueError(f'Не более {MAX_ROWS} строк за один импорт. Разделите таблицу.')
            if any(len(v)>10000 for v in values): raise ValueError(f'Строка {number}: значение превышает 10 000 символов.')
            rows.append({'number':number,'values':values})
        if not rows: raise ValueError('В таблице нет данных после заголовков.')
        return columns,rows
    except ValueError: raise
    except Exception: raise ValueError('Не удалось прочитать таблицу. Проверьте формат и отсутствие пароля на файле.')
    finally:
        if extension=='.xlsx' and 'workbook' in locals(): workbook.close()
        if extension=='.xls' and 'workbook' in locals(): workbook.release_resources()


def suggested_mapping(kind,columns):
    result={}
    for key,label,_ in FIELDS[kind]:
        aliases={norm(v) for v in [key,label,*ALIASES.get(key,[])]}
        for i,column in enumerate(columns):
            if norm(column) in aliases: result[key]=str(i);break
    return result


class UploadImportForm(forms.Form):
    kind=forms.ChoiceField(label='Что импортировать',choices=list(KINDS.items()))
    file=forms.FileField(label='Таблица CSV или Excel',widget=forms.FileInput(attrs={'accept':'.csv,.xlsx,.xls'}))
    sheet=forms.CharField(label='Название листа Excel',required=False,max_length=80,help_text='Если не указано, используется первый лист.')


class MappingForm(forms.Form):
    duplicate_mode=forms.ChoiceField(label='Если запись уже существует',choices=[('skip','Пропустить совпадения'),('update','Обновить совпадения')])
    def __init__(self,*args,batch,**kwargs):
        super().__init__(*args,**kwargs)
        choices=[('','Не загружать')]+[(str(i),name) for i,name in enumerate(batch.columns)]
        for key,label,required in FIELDS[batch.kind]:
            self.fields['map_'+key]=forms.ChoiceField(label=label+(' *' if required else ''),choices=choices,required=required)
    def clean(self):
        data=super().clean();used=[value for key,value in data.items() if key.startswith('map_') and value]
        if len(used)!=len(set(used)): raise forms.ValidationError('Каждую колонку таблицы можно использовать только для одного поля.')
        return data


def decimal(value):
    try:
        number=Decimal(value.replace(' ','').replace('\u00a0','').replace(',','.'))
        if not number.is_finite() or number<0: raise InvalidOperation
        return number
    except (InvalidOperation,AttributeError): raise ValidationError('Укажите неотрицательное число без знака валюты.')


def boolean(value):
    if norm(value) in {'да','yes','true','1','активен'}: return True
    if norm(value) in {'нет','no','false','0','неактивен'}: return False
    raise ValidationError('В поле «Активен» используйте да/нет или 1/0.')


def unique(queryset,message):
    rows=list(queryset[:2])
    if len(rows)>1: raise ValidationError(message)
    return rows[0] if rows else None


def prepare_row(batch,source):
    data={key:source['values'][int(index)].strip() for key,index in batch.mapping.items() if index!='' and source['values'][int(index)].strip()}
    for key,label,required in FIELDS[batch.kind]:
        if required and not data.get(key): raise ValidationError(f'Заполните поле «{label}».')
    values=dict(data);target=None;category_new=None
    if batch.kind=='clients':
        values['phone']=format_phone(data['phone'])
        target=unique(Client.objects.filter(tenant=batch.tenant,normalized_phone=normalize_phone(values['phone'])),'Несколько клиентов с этим телефоном. Сначала объедините дубли в CRM.')
    elif batch.kind=='cars':
        phone=format_phone(data.pop('client_phone'))
        values.pop('client_phone',None)
        client=unique(Client.objects.filter(tenant=batch.tenant,normalized_phone=normalize_phone(phone),archived_at__isnull=True),'Несколько клиентов с этим телефоном.')
        if not client: raise ValidationError('Клиент не найден. Сначала импортируйте клиентов или создайте карточку.')
        values['client_id']=client.pk
        if not data.get('plate') and not data.get('vin'): raise ValidationError('Укажите госномер или VIN для проверки дублей.')
        query=Q(pk__isnull=True)
        if data.get('plate'): query|=Q(normalized_plate=normalize_plate(data['plate']))
        if data.get('vin'): query|=Q(vin__iexact=data['vin']);values['vin']=data['vin'].upper()
        target=unique(Car.objects.filter(tenant=batch.tenant).filter(query),'Госномер и VIN указывают на разные карточки. Исправьте строку.')
        if target and target.client_id!=client.pk: raise ValidationError('Автомобиль уже принадлежит другому клиенту. Проверьте телефон, госномер и VIN.')
        if data.get('vehicle_class'):
            klass=VehicleClass.objects.filter(tenant=batch.tenant,active=True,name__iexact=data['vehicle_class']).first()
            if not klass: raise ValidationError('Класс автомобиля не найден. Создайте его на странице классов.')
            values.pop('vehicle_class');values['vehicle_class_id']=klass.pk
        elif not target:
            klass=VehicleClass.objects.filter(tenant=batch.tenant,active=True,builtin_key='sedan').first()
            if not klass: raise ValidationError('Укажите активный класс автомобиля.')
            values['vehicle_class_id']=klass.pk
    elif batch.kind=='employees':
        if data.get('email'):
            member=batch.tenant.memberships.select_related('user').filter(user__email__iexact=data['email'],status='active').first()
            if not member: raise ValidationError('Этот пользователь ещё не состоит в компании. Сначала создайте ему доступ в CRM.')
            if member.role=='owner': raise ValidationError('Владелец не импортируется как сотрудник.')
            values.pop('email');values['user_id']=member.user_id
        target=unique(Employee.objects.filter(tenant=batch.tenant,location=batch.location).filter(Q(name__iexact=data['name'])|Q(user_id=values.get('user_id')) if values.get('user_id') else Q(name__iexact=data['name'])),'Найдено несколько сотрудников. Уточните имя и email.')
        if target and target.user_id and batch.tenant.memberships.filter(user_id=target.user_id,role='owner').exists(): raise ValidationError('Карточка владельца не изменяется импортом.')
        role=CompanyRole.objects.filter(tenant=batch.tenant,active=True,name__iexact=data['company_role']).first() if data.get('company_role') else None
        if data.get('company_role') and not role: raise ValidationError('Роль не найдена. Создайте её в CRM.')
        if role and role.builtin_key=='owner': raise ValidationError('Роль «Владелец» нельзя назначить импортом.')
        if role: values.pop('company_role');values['company_role_id']=role.pk
        elif not target:
            role=CompanyRole.objects.filter(tenant=batch.tenant,builtin_key='detailer',active=True).first()
            if role: values['company_role_id']=role.pk
        if data.get('salary_mode'):
            choices={norm(label):value for value,label in Employee.SalaryMode.choices}
            mode=choices.get(norm(data['salary_mode']),data['salary_mode'])
            if mode not in Employee.SalaryMode.values: raise ValidationError('Неизвестная схема зарплаты.')
            values['salary_mode']=mode
        for key in ['salary_percent','fixed_monthly_salary']:
            if key in data: values[key]=decimal(data[key])
    else:
        category=InventoryCategory.objects.filter(tenant=batch.tenant,name__iexact=data['category']).first()
        units={'л':'liter','l':'liter','liter':'liter','литр':'liter','мл':'milliliter','ml':'milliliter','milliliter':'milliliter','шт':'piece','шт.':'piece','piece':'piece','pcs':'piece'}
        unit=units.get(data.get('unit','').casefold())
        if data.get('unit') and not unit: raise ValidationError('Единица: л, мл или шт.')
        if category and unit and category.unit!=unit: raise ValidationError('Единица не совпадает с категорией в CRM. Пересчитайте остатки в её единицу.')
        if not category and not unit: raise ValidationError('Для новой категории укажите единицу: л, мл или шт.')
        values.pop('category');values.pop('unit',None)
        if category: values['category_id']=category.pk
        else: category_new={'name':data['category'],'unit':unit}
        target=unique(Consumable.objects.filter(tenant=batch.tenant,location=batch.location,name__iexact=data['name']),'Найдено несколько расходников с этим названием.')
        for key in ['cost_per_unit','planned_stock','actual_stock','minimum_stock']:
            if key in data: values[key]=decimal(data[key])
    if 'active' in data: values['active']=boolean(data['active'])
    model=MODELS[batch.kind]
    obj=target or model(tenant=batch.tenant)
    if batch.kind in {'employees','consumables'}: obj.location=batch.location
    for key,value in values.items(): setattr(obj,key,value)
    if batch.kind=='clients': obj.normalized_phone=normalize_phone(obj.phone)
    obj.full_clean(exclude=['category'] if category_new else [])
    if target and getattr(target,'archived_at',None): raise ValidationError('Карточка архивирована. Восстановите её в CRM перед обновлением.')
    return target,values,category_new


def key_for(kind,values):
    if kind=='clients': return normalize_phone(values['phone'])
    if kind=='cars': return normalize_plate(values.get('plate','')) or values.get('vin','').upper()
    return norm(values['name'])


def preview(batch):
    results=[];seen=set();employee_delta=0
    for row in batch.rows:
        item={'row':row['number'],'label':'','action':'error','message':''}
        try:
            target,values,category=prepare_row(batch,row)
            item['label']=values.get('name') or values.get('plate') or values.get('vin')
            identities={(batch.kind,key_for(batch.kind,values))}
            if batch.kind=='cars':
                identities={(field,str(values[field]).replace(' ','').upper()) for field in ('plate','vin') if values.get(field)}
            if batch.kind=='employees' and values.get('user_id'): identities.add(('user',str(values['user_id'])))
            if identities & seen: raise ValidationError('Повторная строка в файле. Оставьте одну запись с этим идентификатором.')
            seen.update(identities)
            item['action']=batch.duplicate_mode if target else 'create'
            item['target']=str(target.pk) if target else ''
            item['values']=json.loads(json.dumps(values,cls=DjangoJSONEncoder));item['category']=category
            if batch.kind=='employees' and item['action']!='skip':
                was_active=Employee.objects.filter(pk=target.pk,active=True).exists() if target else False
                will_active=values.get('active',target.active if target else True)
                employee_delta+=int(will_active)-int(was_active)
        except ValidationError as exc:
            item['message']=' '.join(f'{key}: {"; ".join(value)}' for key,value in exc.message_dict.items()) if hasattr(exc,'message_dict') else ' '.join(exc.messages)
        results.append(item)
    if batch.kind=='employees':
        limit=active_employee_limit(batch.tenant)
        if limit is not None and counted_employees(batch.tenant).count()+employee_delta>limit:
            for item in results:
                if item['action'] in {'create','update'}:
                    item['action']='error';item['message']='Импорт превышает лимит сотрудников. Отключите лишние карточки или измените подписку.'
    return results


def fingerprint(results):
    return hashlib.sha256(json.dumps(results,sort_keys=True,cls=DjangoJSONEncoder).encode()).hexdigest()


def state(obj):
    return json.loads(json.dumps({f.attname:getattr(obj,f.attname) for f in obj._meta.concrete_fields},cls=DjangoJSONEncoder))


def relations(obj):
    result={}
    for relation in obj._meta.related_objects:
        name=relation.get_accessor_name()
        manager=getattr(obj,name,None)
        if hasattr(manager,'values_list'): result[name]=sorted(str(pk) for pk in manager.values_list('pk',flat=True))
    return result


@transaction.atomic
def apply_import(batch,actor,*,skip_errors=False):
    Tenant.objects.select_for_update().get(pk=batch.tenant_id)
    Subscription.objects.select_for_update().get(tenant=batch.tenant)
    batch=ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status=='applied': return batch
    if batch.status!='preview': raise ValueError('Сначала выполните предварительную проверку.')
    results=preview(batch)
    if fingerprint(results)!=fingerprint(batch.results):
        batch.results=results;batch.save(update_fields=['results'])
        return None
    if any(row['action']=='error' for row in results) and not skip_errors: raise ValueError('Исправьте ошибки или разрешите пропустить строки с ошибками.')
    journal=[];objects=[]
    def remember(kind,obj,before):
        journal.append({'kind':kind,'id':str(obj.pk),'before':before})
        objects.append(obj)
    for item in results:
        if item['action'] not in {'create','update'}: continue
        model=MODELS[batch.kind]
        obj=model.objects.get(tenant=batch.tenant,pk=item['target']) if item['target'] else model(tenant=batch.tenant)
        before=state(obj) if obj.pk and not obj._state.adding else None
        old_plan=getattr(obj,'planned_stock',Decimal('0'));old_actual=getattr(obj,'actual_stock',Decimal('0'))
        values=dict(item['values'])
        if item.get('category'):
            category=InventoryCategory.objects.filter(tenant=batch.tenant,name__iexact=item['category']['name']).first()
            created=category is None
            if created: category=InventoryCategory.objects.create(tenant=batch.tenant,**item['category'])
            if category.unit!=item['category']['unit']: raise ValueError('Единица новой категории не совпадает. Повторите проверку.')
            if created: remember('category',category,None)
            values['category_id']=str(category.pk)
        if batch.kind in {'employees','consumables'}: obj.location=batch.location
        for key,value in values.items(): setattr(obj,key,value)
        if batch.kind=='clients': obj.normalized_phone=normalize_phone(obj.phone)
        obj.full_clean();obj.save();remember(batch.kind,obj,before)
        if batch.kind=='consumables' and (obj.planned_stock!=old_plan or obj.actual_stock!=old_actual):
            movement=InventoryMovement.objects.create(tenant=batch.tenant,location=batch.location,consumable=obj,
                kind=InventoryMovement.Kind.INVENTORY,planned_delta=obj.planned_stock-old_plan,
                actual_delta=obj.actual_stock-old_actual,planned_after=obj.planned_stock,actual_after=obj.actual_stock,
                unit_cost_snapshot=obj.cost_per_unit,created_by=actor,reason=f'Импорт {batch.pk}')
            remember('movement',movement,None)
    for change,obj in zip(journal,objects):
        obj.refresh_from_db();change['after']=state(obj)
        if change['before'] is None: change['relations']=relations(obj)
    batch.changes=journal;batch.results=results;batch.status='applied';batch.applied_at=timezone.now();batch.save()
    log_event(tenant=batch.tenant,actor=actor,action='data_import',entity=batch,
        description=f'Импорт: {KINDS[batch.kind]} ({len(journal)} изменений)')
    return batch


@transaction.atomic
def undo_import(batch,actor):
    Tenant.objects.select_for_update().get(pk=batch.tenant_id)
    batch=ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if batch.status=='undone': return
    if batch.status!='applied': raise ValueError('Этот импорт ещё не применён.')
    objects=[]
    for change in batch.changes:
        obj=MODELS[change['kind']].objects.select_for_update(of=('self',)).filter(tenant=batch.tenant,pk=change['id']).first()
        if not obj or state(obj)!=change['after'] or (change['before'] is None and relations(obj)!=change.get('relations',{})):
            raise ValueError('После импорта карточки изменились или получили новые связи. Автоматическая отмена остановлена, чтобы сохранить новые данные.')
        objects.append(obj)
    for change,obj in reversed(list(zip(batch.changes,objects))):
        if change['before'] is None: obj.delete()
        else:
            restore={key:value for key,value in change['before'].items() if key not in {'id','tenant_id','created_at'}}
            MODELS[change['kind']].objects.filter(pk=obj.pk,tenant=batch.tenant).update(**restore)
    batch.status='undone';batch.undone_at=timezone.now();batch.save(update_fields=['status','undone_at'])
    log_event(tenant=batch.tenant,actor=actor,action='data_import_undo',entity=batch,description='Импорт отменён')
